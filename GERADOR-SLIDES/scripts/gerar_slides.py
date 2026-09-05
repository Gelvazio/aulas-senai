#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Gerador de Slides SENAI
=======================

Le um arquivo Markdown, aplica o padrao visual definido em `padrao_slides.json`
e gera uma apresentacao `.pptx` sobre o template institucional.

Subcomandos
-----------
    template            constroi TEMPLATE-SENAI.pptx a partir do PPTX base
    validar  <arq.md>   confere o Markdown contra as regras do padrao
    gerar    <arq.md>   gera o .pptx final

Padrao de referencia: SLIDE-BASE-EXEMPLO-GELVAZIO-CAMARGO-ITIC.pptx
Documentacao: ../PADRAO-DE-SLIDES.md e ../SINTAXE-MARKDOWN.md
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# O script vive em GERADOR-SLIDES/scripts/; a raiz do gerador e a pasta acima.
SCRIPTS = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(SCRIPTS)
CONFIG = os.path.join(RAIZ, "padrao_slides.json")

# Cores usadas em rodizio nos slides de secao.
CICLO_SECAO = ["primaria", "ciano", "cinza_azul", "ambar", "grafite"]


# --------------------------------------------------------------------------
# Configuracao
# --------------------------------------------------------------------------

def carregar_config(caminho=CONFIG):
    with open(caminho, encoding="utf-8") as fp:
        return json.load(fp)


def cor(cfg, nome):
    """Resolve um nome logico da paleta (ou um hex direto) em RGBColor."""
    hexa = cfg["paleta"].get(nome, nome)
    return RGBColor.from_string(hexa)


# --------------------------------------------------------------------------
# Modelo de dados
# --------------------------------------------------------------------------

class Slide:
    """Um slide ja normalizado, pronto para composicao."""

    def __init__(self, tipo, titulo=""):
        self.tipo = tipo          # capa | secao | conteudo | colunas
        self.titulo = titulo
        self.subtitulo = ""
        self.bullets = []         # [(nivel, texto)]
        self.destaques = []       # [texto]
        self.imagens = []         # [(caminho, legenda)]
        self.tabela = []          # [[celula, ...], ...]
        self.colunas = []         # [{"titulo": str, "bullets": [(nivel, texto)]}]
        self.notas = ""
        self.layout_forcado = ""

    @property
    def vazio(self):
        return not (self.bullets or self.destaques or self.imagens
                    or self.tabela or self.colunas or self.subtitulo)

    def __repr__(self):
        return "<Slide %s %r bullets=%d>" % (self.tipo, self.titulo, len(self.bullets))


# --------------------------------------------------------------------------
# Parser de Markdown
# --------------------------------------------------------------------------

RE_FRONT = re.compile(r"^---\s*$")
RE_H1 = re.compile(r"^#\s+(.*)$")
RE_H2 = re.compile(r"^##\s+(.*)$")
RE_H3 = re.compile(r"^###\s+(.*)$")
RE_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
RE_QUOTE = re.compile(r"^>\s?(.*)$")
RE_IMG = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
RE_TABELA = re.compile(r"^\s*\|(.+)\|\s*$")
RE_SEP_TABELA = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
RE_NOTAS_INI = re.compile(r"^:::\s*notas\s*$", re.I)
RE_NOTAS_FIM = re.compile(r"^:::\s*$")
RE_DIRETIVA = re.compile(r"^<!--\s*layout:\s*([a-z_]+)\s*-->\s*$", re.I)
RE_INLINE = re.compile(r"(\*\*|__|`)")

ROTULOS_CAPA = {
    "uc": "UC",
    "professor": "Professor",
    "carga": "Carga horaria",
    "data": "Data",
}


def limpar_inline(texto):
    """Remove marcacao inline do Markdown — o PPTX aplica o proprio estilo."""
    texto = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", texto)
    texto = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", texto)
    texto = RE_INLINE.sub("", texto)
    return texto.strip()


def parse_front_matter(linhas):
    """Le o bloco `---` ... `---` do topo. Devolve (metadados, indice_corpo)."""
    meta = {}
    if not linhas or not RE_FRONT.match(linhas[0]):
        return meta, 0
    for i in range(1, len(linhas)):
        if RE_FRONT.match(linhas[i]):
            return meta, i + 1
        if ":" in linhas[i]:
            chave, _, valor = linhas[i].partition(":")
            meta[chave.strip().lower()] = valor.strip()
    return meta, 0


def parse_markdown(texto, cfg):
    """Converte o Markdown em (metadados, lista de Slide)."""
    linhas = texto.replace("\r\n", "\n").split("\n")
    meta, inicio = parse_front_matter(linhas)

    slides = []
    estado = {"atual": None, "coluna": None, "tabela": []}
    em_notas = False
    primeiro_h1 = True

    def fechar_tabela():
        if estado["tabela"] and estado["atual"] is not None:
            estado["atual"].tabela = estado["tabela"]
        estado["tabela"] = []

    def novo(tipo, titulo):
        fechar_tabela()
        s = Slide(tipo, limpar_inline(titulo))
        estado["atual"] = s
        estado["coluna"] = None
        slides.append(s)
        return s

    for bruta in linhas[inicio:]:
        linha = bruta.rstrip()

        # --- notas do apresentador ---------------------------------------
        if RE_NOTAS_INI.match(linha):
            em_notas = True
            continue
        if em_notas:
            if RE_NOTAS_FIM.match(linha):
                em_notas = False
            elif estado["atual"] is not None:
                estado["atual"].notas += limpar_inline(linha) + "\n"
            continue

        # --- diretiva de layout -------------------------------------------
        d = RE_DIRETIVA.match(linha)
        if d and estado["atual"] is not None:
            estado["atual"].layout_forcado = d.group(1).lower()
            continue

        # --- titulos -------------------------------------------------------
        m = RE_H1.match(linha)
        if m:
            novo("capa" if primeiro_h1 else "secao", m.group(1))
            primeiro_h1 = False
            continue

        m = RE_H2.match(linha)
        if m:
            novo("conteudo", m.group(1))
            continue

        m = RE_H3.match(linha)
        if m:
            if estado["atual"] is None:
                novo("conteudo", "")
            fechar_tabela()
            estado["atual"].tipo = "colunas"
            estado["coluna"] = {"titulo": limpar_inline(m.group(1)), "bullets": []}
            estado["atual"].colunas.append(estado["coluna"])
            continue

        if estado["atual"] is None:
            if linha.strip():
                novo("conteudo", "")
            else:
                continue

        atual = estado["atual"]

        # --- tabela ---------------------------------------------------------
        if RE_TABELA.match(linha):
            if not RE_SEP_TABELA.match(linha):
                celulas = [limpar_inline(c) for c in linha.strip().strip("|").split("|")]
                estado["tabela"].append(celulas)
            continue
        if estado["tabela"]:
            fechar_tabela()

        # --- imagem ----------------------------------------------------------
        m = RE_IMG.match(linha.strip())
        if m:
            atual.imagens.append((m.group(2).strip(), limpar_inline(m.group(1))))
            continue

        # --- destaque --------------------------------------------------------
        m = RE_QUOTE.match(linha)
        if m:
            conteudo = limpar_inline(m.group(1))
            if conteudo:
                atual.destaques.append(conteudo)
            continue

        # --- bullet -----------------------------------------------------------
        m = RE_BULLET.match(linha)
        if m:
            recuo = len(m.group(1).replace("\t", "  "))
            nivel = min(recuo // 2, cfg["limites"]["max_niveis_indentacao"] - 1)
            conteudo = limpar_inline(m.group(2))
            if not conteudo:
                continue
            if estado["coluna"] is not None:
                estado["coluna"]["bullets"].append((nivel, conteudo))
            else:
                atual.bullets.append((nivel, conteudo))
            continue

        # --- paragrafo solto ---------------------------------------------------
        conteudo = limpar_inline(linha)
        if not conteudo:
            continue
        if atual.tipo == "capa" and not atual.subtitulo:
            atual.subtitulo = conteudo
        elif estado["coluna"] is not None:
            estado["coluna"]["bullets"].append((0, conteudo))
        else:
            atual.bullets.append((0, conteudo))

    fechar_tabela()
    return meta, [s for s in slides if s.titulo or not s.vazio]


# --------------------------------------------------------------------------
# Normalizacao — quebra automatica de slides sobrecarregados
# --------------------------------------------------------------------------

def dividir_excedentes(slides, cfg):
    """Quebra em varios slides todo slide que passa do maximo de bullets."""
    limite = cfg["limites"]["max_bullets_por_slide"]
    saida = []

    for s in slides:
        if s.tipo in ("capa", "secao", "colunas") or len(s.bullets) <= limite:
            saida.append(s)
            continue

        blocos = [s.bullets[i:i + limite] for i in range(0, len(s.bullets), limite)]
        for i, bloco in enumerate(blocos):
            parte = Slide("conteudo", s.titulo if i == 0 else s.titulo + " (cont.)")
            parte.bullets = bloco
            parte.layout_forcado = s.layout_forcado
            if i == len(blocos) - 1:
                parte.destaques = s.destaques
                parte.imagens = s.imagens
                parte.tabela = s.tabela
                parte.notas = s.notas
            saida.append(parte)

    return saida


# --------------------------------------------------------------------------
# Composicao do PPTX
# --------------------------------------------------------------------------

class Compositor:
    """Desenha os slides aplicando a grade, a tipografia e a paleta do padrao."""

    def __init__(self, cfg, template):
        self.cfg = cfg
        self.prs = Presentation(template)
        self.g = cfg["grade"]
        self.tipo = cfg["tipografia"]
        self.familia = self.tipo["familia"]

    # -- utilitarios --------------------------------------------------------

    def layout(self, nome_logico):
        alvo = self.cfg["layouts"].get(nome_logico, "TITLE_ONLY")
        for lay in self.prs.slide_layouts:
            if lay.name == alvo:
                return lay
        for lay in self.prs.slide_layouts:
            if lay.name == "TITLE_ONLY":
                return lay
        return self.prs.slide_layouts[0]

    def estilo(self, run, chave):
        e = self.tipo[chave]
        run.font.name = self.familia
        run.font.size = Pt(e["pt"])
        run.font.bold = e["bold"]
        run.font.color.rgb = RGBColor.from_string(e["cor"])

    @staticmethod
    def remover(shape):
        shape._element.getparent().remove(shape._element)

    def limpar_placeholders(self, slide, manter_titulo=True):
        titulo = slide.shapes.title if manter_titulo else None
        for shape in list(slide.shapes):
            if shape.is_placeholder and shape is not titulo:
                self.remover(shape)

    def caixa(self, slide, esq, topo, larg, alt):
        cx = slide.shapes.add_textbox(esq, topo, larg, alt)
        tf = cx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        return cx, tf

    def largura_util(self):
        return self.prs.slide_width - Inches(
            self.g["margem_esquerda_pol"] + self.g["margem_direita_pol"])

    def area_corpo(self):
        esq = Inches(self.g["margem_esquerda_pol"])
        topo = Inches(self.g["topo_corpo_pol"])
        return esq, topo, self.largura_util(), Inches(self.g["altura_corpo_pol"])

    # -- blocos reutilizaveis ------------------------------------------------

    def por_titulo(self, slide, texto):
        alvo = slide.shapes.title
        if alvo is None:
            _, tf = self.caixa(slide,
                               Inches(self.g["margem_esquerda_pol"]),
                               Inches(self.g["margem_topo_pol"]),
                               self.largura_util(), Inches(0.95))
        else:
            tf = alvo.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ""
        run = p.add_run()
        run.text = texto
        self.estilo(run, "titulo_slide")

    def por_bullets(self, tf, itens, usar_primeiro=True):
        for i, (nivel, texto) in enumerate(itens):
            p = tf.paragraphs[0] if (i == 0 and usar_primeiro) else tf.add_paragraph()
            p.level = nivel
            p.space_after = Pt(8 if nivel == 0 else 4)
            if nivel == 0:
                marcador = ""
            elif nivel == 1:
                marcador = "• "
            else:
                marcador = "◦ "
            run = p.add_run()
            run.text = marcador + texto
            self.estilo(run, "bullet_n%d" % min(nivel + 1, 3))

    def por_tabela(self, slide, dados, esq, topo, larg):
        lim = self.cfg["limites"]
        linhas = dados[:lim["max_linhas_tabela"] + 1]
        n_col = min(max(len(l) for l in linhas), lim["max_colunas_tabela"])
        altura = Inches(0.34) * len(linhas)
        tbl = slide.shapes.add_table(len(linhas), n_col, esq, topo, larg, altura).table

        for i, linha in enumerate(linhas):
            for j in range(n_col):
                cel = tbl.cell(i, j)
                cel.text = linha[j] if j < len(linha) else ""
                cel.margin_left = Inches(0.08)
                cel.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = cel.text_frame.paragraphs[0]
                if not p.runs:
                    p.add_run().text = ""
                for run in p.runs:
                    self.estilo(run, "tabela_cabecalho" if i == 0 else "tabela_celula")
        return topo + altura + Inches(0.12)

    def por_imagem(self, slide, caminho, legenda, esq, topo, larg, alt_max):
        if not os.path.isabs(caminho):
            caminho = os.path.normpath(os.path.join(RAIZ, caminho))
        if not os.path.exists(caminho):
            print("  [aviso] imagem nao encontrada, ignorada: %s" % caminho)
            return topo

        alt_max = max(int(alt_max), int(Inches(1.0)))
        pic = slide.shapes.add_picture(caminho, esq, topo)
        escala = min(larg / pic.width, alt_max / pic.height, 1.0)
        pic.width = int(pic.width * escala)
        pic.height = int(pic.height * escala)
        pic.left = int(esq + (larg - pic.width) / 2)
        pic.top = int(topo)

        cursor = pic.top + pic.height + Inches(0.06)
        if legenda:
            _, tf = self.caixa(slide, esq, cursor, larg, Inches(0.28))
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = legenda
            self.estilo(run, "legenda")
            cursor += Inches(0.3)
        return cursor

    def por_destaque(self, slide, texto, esq, topo, larg):
        altura = Inches(0.62)
        cx = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, esq, topo, larg, altura)
        cx.fill.solid()
        cx.fill.fore_color.rgb = cor(self.cfg, "fundo_suave")
        cx.line.color.rgb = cor(self.cfg, "ambar")
        cx.line.width = Pt(1.5)
        cx.shadow.inherit = False
        tf = cx.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = ""
        run = p.add_run()
        run.text = texto
        self.estilo(run, "destaque")
        return topo + altura + Inches(0.12)

    def por_rodape(self, slide, meta, numero):
        cfg_rod = self.cfg["rodape"]
        if not cfg_rod["ativo"]:
            return
        texto = cfg_rod["modelo"].format(
            uc=meta.get("uc", meta.get("titulo", "")),
            professor=meta.get("professor", ""),
            n=numero,
        )
        topo = self.prs.slide_height - Inches(self.g["margem_rodape_pol"])
        _, tf = self.caixa(slide, Inches(self.g["margem_esquerda_pol"]),
                           topo, self.largura_util(), Inches(0.22))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = texto
        self.estilo(run, "rodape")

    # -- tipos de slide -------------------------------------------------------

    def capa(self, s, meta):
        slide = self.prs.slides.add_slide(self.layout("capa"))
        self.limpar_placeholders(slide, manter_titulo=False)

        esq = Inches(self.g["margem_esquerda_pol"])
        larg = self.largura_util()

        faixa = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                       Inches(0.18), self.prs.slide_height)
        faixa.fill.solid()
        faixa.fill.fore_color.rgb = cor(self.cfg, "primaria")
        faixa.line.fill.background()
        faixa.shadow.inherit = False

        _, tf = self.caixa(slide, esq, Inches(1.55), larg, Inches(1.4))
        run = tf.paragraphs[0].add_run()
        run.text = s.titulo
        self.estilo(run, "titulo_capa")

        linhas = [s.subtitulo] if s.subtitulo else []
        for chave, rotulo in ROTULOS_CAPA.items():
            if meta.get(chave):
                linhas.append("%s: %s" % (rotulo, meta[chave]))
        if linhas:
            _, tf2 = self.caixa(slide, esq, Inches(3.15), larg, Inches(1.9))
            for i, texto in enumerate(linhas):
                p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
                p.space_after = Pt(4)
                run = p.add_run()
                run.text = texto
                self.estilo(run, "subtitulo_capa")
        return slide

    def secao(self, s, indice):
        slide = self.prs.slides.add_slide(self.layout("secao"))
        self.limpar_placeholders(slide, manter_titulo=False)

        fundo = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                       self.prs.slide_width, self.prs.slide_height)
        fundo.fill.solid()
        fundo.fill.fore_color.rgb = cor(self.cfg, CICLO_SECAO[indice % len(CICLO_SECAO)])
        fundo.line.fill.background()
        fundo.shadow.inherit = False

        _, tf = self.caixa(slide, Inches(self.g["margem_esquerda_pol"]),
                           Inches(2.05), self.largura_util(), Inches(1.6))
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = tf.paragraphs[0].add_run()
        run.text = s.titulo
        self.estilo(run, "titulo_secao")
        return slide

    def colunas(self, s):
        slide = self.prs.slides.add_slide(self.layout("livre"))
        self.limpar_placeholders(slide)
        self.por_titulo(slide, s.titulo)

        esq, topo, larg, alt = self.area_corpo()
        n = max(len(s.colunas), 1)
        vao = Inches(self.g["espaco_entre_colunas_pol"])
        larg_col = int((larg - vao * (n - 1)) / n)

        for i, col in enumerate(s.colunas):
            x = int(esq + i * (larg_col + vao))
            _, tf = self.caixa(slide, x, topo, larg_col, alt)
            p = tf.paragraphs[0]
            p.space_after = Pt(10)
            run = p.add_run()
            run.text = col["titulo"]
            self.estilo(run, "subtitulo")
            if col["bullets"]:
                self.por_bullets(tf, col["bullets"], usar_primeiro=False)
        return slide

    def conteudo(self, s):
        slide = self.prs.slides.add_slide(self.layout("livre"))
        self.limpar_placeholders(slide)
        self.por_titulo(slide, s.titulo)

        esq, topo, larg, alt = self.area_corpo()
        cursor = topo

        if s.bullets:
            altura = min(int(alt), int(Inches(0.42) * len(s.bullets) + Inches(0.2)))
            _, tf = self.caixa(slide, esq, cursor, larg, altura)
            self.por_bullets(tf, s.bullets)
            cursor += altura + Inches(0.12)

        for texto in s.destaques:
            cursor = self.por_destaque(slide, texto, esq, cursor, larg)

        if s.tabela:
            cursor = self.por_tabela(slide, s.tabela, esq, cursor, larg)

        for caminho, legenda in s.imagens:
            restante = topo + alt - cursor
            cursor = self.por_imagem(slide, caminho, legenda, esq, cursor, larg, restante)

        return slide

    # -- montagem --------------------------------------------------------------

    def montar(self, meta, slides, saida):
        n_secao = 0
        for i, s in enumerate(slides):
            if s.tipo == "capa":
                slide = self.capa(s, meta)
            elif s.tipo == "secao":
                slide = self.secao(s, n_secao)
                n_secao += 1
            elif s.tipo == "colunas":
                slide = self.colunas(s)
            else:
                slide = self.conteudo(s)

            if s.notas.strip():
                slide.notes_slide.notes_text_frame.text = s.notas.strip()
            if i > 0 or self.cfg["rodape"]["numerar_capa"]:
                self.por_rodape(slide, meta, i + 1)

        destino = os.path.dirname(os.path.abspath(saida))
        if destino:
            os.makedirs(destino, exist_ok=True)
        self.prs.save(saida)
        return len(slides)


# --------------------------------------------------------------------------
# Validacao
# --------------------------------------------------------------------------

def validar(meta, slides, cfg):
    """Devolve a lista de problemas. Linhas iniciadas por ERRO bloqueiam a geracao."""
    lim = cfg["limites"]
    problemas = []

    if lim["min_slides_por_deck"] is not None and len(slides) < lim["min_slides_por_deck"]:
        problemas.append(
            "ERRO  deck com %d slides; o padrao exige no minimo %d"
            % (len(slides), lim["min_slides_por_deck"]))

    if not slides or slides[0].tipo != "capa":
        problemas.append("ERRO  o deck deve comecar por uma capa (linha '# Titulo')")

    for chave in ("uc", "professor"):
        if not meta.get(chave):
            problemas.append("AVISO front-matter sem o campo '%s'" % chave)

    for i, s in enumerate(slides, start=1):
        marca = "slide %d (%s)" % (i, s.titulo[:40] or "sem titulo")
        if len(s.titulo) > lim["max_caracteres_titulo"]:
            problemas.append(
                "AVISO %s: titulo com %d caracteres (maximo %d)"
                % (marca, len(s.titulo), lim["max_caracteres_titulo"]))
        if len(s.bullets) > lim["max_bullets_por_slide"]:
            problemas.append(
                "AVISO %s: %d bullets; sera quebrado automaticamente em %d por slide"
                % (marca, len(s.bullets), lim["max_bullets_por_slide"]))
        for _, texto in s.bullets:
            if len(texto) > lim["max_caracteres_por_bullet"]:
                problemas.append(
                    "AVISO %s: bullet com %d caracteres — encurte para ate %d"
                    % (marca, len(texto), lim["max_caracteres_por_bullet"]))
        for caminho, _ in s.imagens:
            alvo = caminho if os.path.isabs(caminho) else os.path.join(RAIZ, caminho)
            if not os.path.exists(alvo):
                problemas.append("ERRO  %s: imagem inexistente — %s" % (marca, caminho))

    return problemas


# --------------------------------------------------------------------------
# Template enxuto
# --------------------------------------------------------------------------

def construir_template(cfg):
    """Remove os slides do PPTX base, preservando tema, master, layouts e fontes."""
    base = os.path.join(RAIZ, cfg["apresentacao"]["base"])
    destino = os.path.join(RAIZ, cfg["apresentacao"]["template"])
    if not os.path.exists(base):
        raise SystemExit("PPTX base nao encontrado: %s" % base)

    prs = Presentation(base)
    lst = prs.slides._sldIdLst
    ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    removidos = 0
    for elemento in list(lst):
        prs.part.drop_rel(elemento.get(ns))
        lst.remove(elemento)
        removidos += 1
    prs.save(destino)

    mb = os.path.getsize(destino) / 1048576.0
    print("[ok] template criado: %s" % destino)
    print("     %d slides removidos | %d layouts preservados | %.2f MB"
          % (removidos, len(prs.slide_layouts), mb))
    return destino


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def ler(caminho):
    with open(caminho, encoding="utf-8") as fp:
        return fp.read()


def cmd_validar(args):
    cfg = carregar_config()
    meta, slides = parse_markdown(ler(args.entrada), cfg)
    problemas = validar(meta, slides, cfg)

    tipos = {}
    for s in slides:
        tipos[s.tipo] = tipos.get(s.tipo, 0) + 1

    print("Arquivo : %s" % args.entrada)
    print("Slides  : %d" % len(slides))
    print("Tipos   : %s" % tipos)

    if not problemas:
        print("[ok] nenhum problema encontrado")
        return 0

    print("\n%d ponto(s) de atencao:" % len(problemas))
    for p in problemas:
        print("  " + p)
    return 1 if any(p.startswith("ERRO") for p in problemas) else 0


def cmd_gerar(args):
    cfg = carregar_config()
    template = args.template or os.path.join(RAIZ, cfg["apresentacao"]["template"])
    if not os.path.exists(template):
        print("[info] template ausente — construindo a partir do PPTX base")
        template = construir_template(cfg)

    meta, slides = parse_markdown(ler(args.entrada), cfg)
    for chave in ("titulo", "uc", "professor", "carga", "data"):
        if getattr(args, chave, None):
            meta[chave] = getattr(args, chave)

    problemas = validar(meta, slides, cfg)
    for p in problemas:
        print("  " + p)
    if any(p.startswith("ERRO") for p in problemas) and not args.forcar:
        print("\n[falhou] corrija os ERROs acima ou use --forcar")
        return 1

    slides = dividir_excedentes(slides, cfg)
    saida = args.saida or os.path.join(
        RAIZ, "SAIDA",
        os.path.splitext(os.path.basename(args.entrada))[0] + ".pptx")

    total = Compositor(cfg, template).montar(meta, slides, saida)
    mb = os.path.getsize(saida) / 1048576.0
    print("\n[ok] %d slides gerados" % total)
    print("     %s  (%.2f MB)" % (saida, mb))
    return 0


def cmd_template(args):
    construir_template(carregar_config())
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(
        prog="gerar_slides.py",
        description="Gerador de slides SENAI — Markdown para PPTX no padrao institucional.")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("template", help="constroi TEMPLATE-SENAI.pptx a partir do PPTX base")
    p.set_defaults(func=cmd_template)

    p = sub.add_parser("validar", help="confere o Markdown contra as regras do padrao")
    p.add_argument("entrada", help="arquivo .md de entrada")
    p.set_defaults(func=cmd_validar)

    p = sub.add_parser("gerar", help="gera o .pptx final")
    p.add_argument("entrada", help="arquivo .md de entrada")
    p.add_argument("-o", "--saida", help="caminho do .pptx de saida")
    p.add_argument("-t", "--template", help="template alternativo")
    p.add_argument("--titulo", help="sobrescreve o titulo do front-matter")
    p.add_argument("--uc", help="sobrescreve a UC do front-matter")
    p.add_argument("--professor", help="sobrescreve o professor do front-matter")
    p.add_argument("--carga", help="sobrescreve a carga horaria")
    p.add_argument("--data", help="sobrescreve a data")
    p.add_argument("--forcar", action="store_true", help="gera mesmo com ERROs de validacao")
    p.set_defaults(func=cmd_gerar)

    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
