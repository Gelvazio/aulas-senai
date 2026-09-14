#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pathlib import Path

SENAI_BLUE = RGBColor(0, 67, 132)
SENAI_ORANGE = RGBColor(247, 149, 29)
DARK_GRAY = RGBColor(60, 60, 60)

class ExcelConsolidator:
    def __init__(self, base_path):
        self.base_path = Path(base_path)
        self.aulas_folder = self.base_path / "AULAS-CHALKIE-AI-COLORIDA"
        self.slides = []

    def extract_pptx(self, pptx_file, aula_num):
        print(f"[*] Lendo Aula {aula_num}: {pptx_file.name}")
        try:
            prs = Presentation(pptx_file)
            slides_info = []
            for idx, slide in enumerate(prs.slides):
                texts = []
                title = ""
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        if hasattr(shape, "name") and "Title" in shape.name:
                            title = shape.text.strip()
                        else:
                            texts.append(shape.text.strip())
                if title or texts:
                    slides_info.append({
                        'titulo': title,
                        'conteudo': texts,
                        'aula': aula_num
                    })
            print(f"    [OK] {len(slides_info)} slides extraidos")
            return slides_info
        except Exception as e:
            print(f"    [ERR] {e}")
            return []

    def consolidate(self):
        print("\n[*] Consolidando slides...")

        all_slides = []
        # Descobrir arquivos automaticamente
        for pptx_file in sorted(self.aulas_folder.glob("*.pptx")):
            # Extrair numero da aula do nome do arquivo
            filename = pptx_file.name
            if filename.startswith(("3-", "4-", "5-", "6-")):
                aula_num = int(filename[0])
                slides = self.extract_pptx(pptx_file, aula_num)
                all_slides.extend(slides)

        print(f"\n[*] Total de slides extraidos: {len(all_slides)}")

        consolidated = []
        consolidated.append(self._make_slide("EXCEL - Consolidado", "Analise de Dados", ""))

        consolidated.append(self._make_slide(
            "Indice",
            "Conteudo",
            "1. Basics: Interface, Formulas, Funcoes\n"
            "2. Intermediario: Formatacao, Validacao\n"
            "3. Avancado: PROCV, Tabelas Dinamicas\n"
            "4. Visualizacao: Graficos e Dashboards"
        ))

        basic = [s for s in all_slides if s['aula'] in [3, 4]]
        adv = [s for s in all_slides if s['aula'] in [5, 6]]

        consolidated.append(self._make_slide("Excel Basico", f"{len(basic)} slides",
            "Interface e Estrutura\nFormulas Basicas\nFuncoes Essenciais\nFormatacao"))

        consolidated.append(self._make_slide("Excel Avancado", f"{len(adv)} slides",
            "Funcoes de Busca (PROCV)\nFuncoes Condicionais\nTabelas Dinamicas\nGraficos"))

        for slide in basic[:10]:
            consolidated.append(self._make_slide(
                slide['titulo'][:40],
                "Basico",
                "\n".join(slide['conteudo'][:100] if slide['conteudo'] else [""])
            ))

        for slide in adv[:15]:
            consolidated.append(self._make_slide(
                slide['titulo'][:40],
                "Avancado",
                "\n".join(slide['conteudo'][:100] if slide['conteudo'] else [""])
            ))

        consolidated.append(self._make_slide(
            "Conclusao",
            "Proximos Passos",
            "1. Pratica com dados reais\n"
            "2. Criacao de dashboards\n"
            "3. Automacao com macros\n"
            "4. Analise de dados"
        ))

        self.slides = consolidated[:50]
        print(f"[OK] {len(self.slides)} slides consolidados")
        return self.slides

    def _make_slide(self, title, subtitle, content):
        return {'title': title, 'subtitle': subtitle, 'content': content}

    def create_pptx(self, output):
        print(f"\n[*] Criando {output}...")
        prs = Presentation()
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)

        for idx, slide_data in enumerate(self.slides, 1):
            blank = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank)

            bg = slide.background
            fill = bg.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor(255, 255, 255)

            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
            tf = title_box.text_frame
            tf.text = slide_data['title']
            p = tf.paragraphs[0]
            p.font.size = Pt(44)
            p.font.bold = True
            p.font.color.rgb = SENAI_BLUE

            if slide_data['subtitle']:
                sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(9), Inches(0.6))
                sf = sub_box.text_frame
                sf.text = slide_data['subtitle']
                sp = sf.paragraphs[0]
                sp.font.size = Pt(18)
                sp.font.color.rgb = SENAI_ORANGE

            if slide_data['content']:
                content_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(9), Inches(4.8))
                cf = content_box.text_frame
                cf.word_wrap = True
                cf.text = slide_data['content']
                for cp in cf.paragraphs:
                    cp.font.size = Pt(14)
                    cp.font.color.rgb = DARK_GRAY

            footer_box = slide.shapes.add_textbox(Inches(9), Inches(7.2), Inches(0.8), Inches(0.3))
            ff = footer_box.text_frame
            ff.text = str(idx)
            fp = ff.paragraphs[0]
            fp.font.size = Pt(10)
            fp.font.color.rgb = RGBColor(200, 200, 200)
            fp.alignment = PP_ALIGN.RIGHT

            line = slide.shapes.add_shape(1, Inches(0.5), Inches(0.1), Inches(9), Inches(0))
            line.line.color.rgb = SENAI_ORANGE
            line.line.width = Pt(3)

        prs.save(output)
        print(f"[OK] Arquivo criado: {output}")
        print(f"[+] Slides totais: {len(self.slides)}/50")

    def run(self):
        print("=" * 60)
        print("[*] CONSOLIDADOR DE SLIDES EXCEL")
        print("=" * 60)
        self.consolidate()
        output = str(self.base_path / "EXCEL.pptx")
        self.create_pptx(output)
        print("=" * 60)
        print("[OK] CONCLUIDO COM SUCESSO!")
        print("=" * 60)

if __name__ == "__main__":
    base = Path(__file__).parent / "MATERIAIS" / "GESTAO_E_CONTROLE_MATERIAIS" / "ANALISE_DADOS_APLICADA_GESTAO"
    if not base.exists():
        print(f"[ERR] Pasta nao encontrada: {base}")
        sys.exit(1)
    try:
        from pptx import Presentation
    except ImportError:
        print("[ERR] python-pptx nao instalado. Execute: pip install python-pptx")
        sys.exit(1)
    consolidator = ExcelConsolidator(base)
    consolidator.run()
