# -*- coding: utf-8 -*-
"""
Marcar APENAS com ✅ a alternativa correta
SEM adicionar texto "RESPOSTA CORRETA"
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PASTA = r"C:\fontes\aulas-senai\MATERIAIS\GESTAO_E_CONTROLE_MATERIAIS\ANALISE_DADOS_APLICADA_GESTAO\ATIVIDADES"

GABARITOS = {
    "GABARITO-ATIVIDADE-01-ESTATISTICA-E-PROGRESSOES.docx": "E A E E A A A A A A A A",
    "GABARITO-ATIVIDADE-02-CONCEITOS-FUNDAMENTOS-EXCEL.docx": "E A A A E A A A A A A A",
    "GABARITO-ATIVIDADE-03-FUNCOES-DE-BUSCA-AVANCADAS.docx": "A A A A A A A A A A A A A",
    "GABARITO-ATIVIDADE-04-DESIGN-DASHBOARD-E-KPIS.docx": "A A A A A A A A A A A A A A",
}

def marcar_gabarito_limpo(docx_path, respostas_str):
    """Marca APENAS com ✅ a alternativa correta (sem texto extra)"""

    respostas = respostas_str.strip().split()
    doc = Document(docx_path)

    questao_num = 0

    # Processar TODAS as tabelas (exceto cabecalho 0)
    for t_idx, table in enumerate(doc.tables):
        if t_idx == 0:
            continue

        questao_num += 1

        if questao_num <= len(respostas):
            resposta_correta = respostas[questao_num - 1].lower()

            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        texto = para.text.strip()

                        # Procurar apenas por "X)" onde X eh a resposta correta
                        if texto.startswith(resposta_correta + ")"):
                            # Reescrever APENAS com ✅ na frente
                            para.clear()

                            # Adicionar checkmark em verde + alternativa + texto
                            run_check = para.add_run("✅ ")
                            run_check.font.color.rgb = RGBColor(0, 128, 0)
                            run_check.bold = True

                            run_rest = para.add_run(texto)
                            run_rest.font.size = Pt(11)

    # Adicionar tabela de gabarito
    doc.add_paragraph()
    doc.add_paragraph()

    titulo = doc.add_paragraph("GABARITO COMPLETO")
    titulo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.runs[0].bold = True
    titulo.runs[0].font.size = Pt(14)

    tabela = doc.add_table(rows=2, cols=len(respostas))
    tabela.style = 'Light Grid Accent 1'

    for i in range(len(respostas)):
        tabela.rows[0].cells[i].text = f"Q{i+1}"
        tabela.rows[0].cells[i].paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tabela.rows[0].cells[i].paragraphs[0].runs[0].bold = True

    for i, resp in enumerate(respostas):
        celula = tabela.rows[1].cells[i]
        celula.text = f"✅ {resp.upper()}"
        para = celula.paragraphs[0]
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.runs[0].bold = True
        para.runs[0].font.color.rgb = RGBColor(0, 128, 0)

    doc.save(docx_path)
    print(f"  OK")

def main():
    print("GABARITO LIMPO - APENAS CHECKMARK SEM TEXTO")
    print("=" * 60)

    for gabarito, respostas in GABARITOS.items():
        print(f"{gabarito}")
        try:
            marcar_gabarito_limpo(f"{PASTA}\\{gabarito}", respostas)
        except Exception as e:
            print(f"  ERRO: {e}")

    print("=" * 60)
    print("Concluido!")

if __name__ == "__main__":
    main()
