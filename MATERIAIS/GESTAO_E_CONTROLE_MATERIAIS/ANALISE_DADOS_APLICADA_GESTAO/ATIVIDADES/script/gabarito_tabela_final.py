# -*- coding: utf-8 -*-
"""
Adicionar APENAS tabela de gabarito no final dos gabaritos
SEM marcar nada nas questoes individuais
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PASTA = r"C:\fontes\aulas-senai\MATERIAIS\GESTAO_E_CONTROLE_MATERIAIS\ANALISE_DADOS_APLICADA_GESTAO\ATIVIDADES"

GABARITOS = {
    "GABARITO-ATIVIDADE-01-ESTATISTICA-E-PROGRESSOES.docx": "E A E E A A A A A A A A",
    "GABARITO-ATIVIDADE-02-CONCEITOS-FUNDAMENTOS-EXCEL.docx": "E A A A E A A A A A A A",
    "GABARITO-ATIVIDADE-03-FUNCOES-DE-BUSCA-AVANCADAS.docx": "A A A A A A A A A A A A A",
    "GABARITO-ATIVIDADE-04-DESIGN-DASHBOARD-E-KPIS.docx": "A A A A A A A A A A A A A A",
}

def adicionar_tabela_gabarito(docx_path, respostas_str):
    """Adiciona APENAS tabela de gabarito no final"""

    respostas = respostas_str.strip().split()
    doc = Document(docx_path)

    # Adicionar paragrafos de espacamento
    doc.add_paragraph()
    doc.add_paragraph()

    # Titulo
    titulo = doc.add_paragraph("GABARITO COMPLETO")
    titulo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_run = titulo.runs[0]
    titulo_run.bold = True
    titulo_run.font.size = Pt(14)

    # Tabela de gabarito
    tabela_gab = doc.add_table(rows=2, cols=len(respostas))
    tabela_gab.style = 'Light Grid Accent 1'

    # Header - Questoes
    for i in range(len(respostas)):
        celula = tabela_gab.rows[0].cells[i]
        celula.text = f"Q{i+1}"
        for para in celula.paragraphs:
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.runs[0].bold = True

    # Respostas (sem cores, apenas clean)
    for i, resp in enumerate(respostas):
        celula = tabela_gab.rows[1].cells[i]
        celula.text = resp.upper()
        for para in celula.paragraphs:
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.runs[0].bold = True

    doc.save(docx_path)
    print(f"  OK - Tabela adicionada ({len(respostas)} questoes)")

def main():
    print("ADICIONAR TABELA DE GABARITO NO FINAL")
    print("=" * 60)

    for gabarito_path, respostas in GABARITOS.items():
        caminho = f"{PASTA}\\{gabarito_path}"
        try:
            print(f"\n{gabarito_path}")
            adicionar_tabela_gabarito(caminho, respostas)
        except Exception as e:
            print(f"  ERRO: {e}")

    print("\n" + "=" * 60)
    print("Concluido!")

if __name__ == "__main__":
    main()
