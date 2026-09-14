# -*- coding: utf-8 -*-
"""
Marcar APENAS a alternativa correta em cada questao
Adicionar tabela de gabarito no final
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

PASTA = r"C:\fontes\aulas-senai\MATERIAIS\GESTAO_E_CONTROLE_MATERIAIS\ANALISE_DADOS_APLICADA_GESTAO\ATIVIDADES"

GABARITOS = {
    "GABARITO-ATIVIDADE-01-ESTATISTICA-E-PROGRESSOES.docx": "E A E E A A A A A A A A",
    "GABARITO-ATIVIDADE-02-CONCEITOS-FUNDAMENTOS-EXCEL.docx": "E A A A E A A A A A A A",
    "GABARITO-ATIVIDADE-03-FUNCOES-DE-BUSCA-AVANCADAS.docx": "A A A A A A A A A A A A A",
    "GABARITO-ATIVIDADE-04-DESIGN-DASHBOARD-E-KPIS.docx": "A A A A A A A A A A A A A A",
}

def marcar_gabarito(docx_path, respostas_str):
    """Marca APENAS a alternativa correta de cada questao"""

    respostas = respostas_str.strip().split()
    doc = Document(docx_path)

    questao_num = 0

    # Processar TODAS as tabelas (exceto cabecalho 0)
    for t_idx, table in enumerate(doc.tables):
        if t_idx == 0:
            continue

        questao_num += 1

        if questao_num <= len(respostas):
            resposta_correta = respostas[questao_num - 1].lower()  # a, b, c, d, e

            # Processar cada celula da tabela
            for row in table.rows:
                for cell in row.cells:
                    # Processar cada paragrafo
                    for para in cell.paragraphs:
                        texto = para.text.strip()

                        # Verificar se eh a alternativa correta (ex: "a)" ou "a) ")
                        if texto.startswith(resposta_correta + ")"):
                            # Eh a alternativa correta
                            # Limpar e reescrever com marca
                            para.clear()

                            # Adicionar alternativa com marca
                            run = para.add_run(f"{resposta_correta.upper()}) ")
                            run.font.size = Pt(11)

                            # Adicionar resto do texto
                            resto_texto = texto[len(resposta_correta)+1:].strip()
                            if resto_texto:
                                run2 = para.add_run(resto_texto)
                                run2.font.size = Pt(11)

                            # Adicionar marca de resposta correta
                            run3 = para.add_run("  ✅")
                            run3.font.color.rgb = RGBColor(0, 128, 0)  # Verde
                            run3.bold = True

    # Adicionar TABELA DE GABARITO no final
    doc.add_paragraph()
    doc.add_paragraph()

    titulo = doc.add_paragraph("GABARITO COMPLETO")
    titulo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_run = titulo.runs[0]
    titulo_run.bold = True
    titulo_run.font.size = Pt(14)

    # Criar tabela
    tabela_gab = doc.add_table(rows=2, cols=len(respostas))
    tabela_gab.style = 'Light Grid Accent 1'

    # Header
    for i in range(len(respostas)):
        celula = tabela_gab.rows[0].cells[i]
        celula.text = f"Q{i+1}"
        for para in celula.paragraphs:
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.runs[0].bold = True

    # Respostas
    for i, resp in enumerate(respostas):
        celula = tabela_gab.rows[1].cells[i]
        celula.text = f"✅ {resp.upper()}"
        for para in celula.paragraphs:
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.runs[0].bold = True
            para.runs[0].font.color.rgb = RGBColor(0, 128, 0)  # Verde

    doc.save(docx_path)
    print(f"  OK - Marcado: {len(respostas)} questoes + tabela")

def main():
    print("MARCAR GABARITO - APENAS ALTERNATIVA CORRETA")
    print("=" * 60)

    for gabarito_path, respostas in GABARITOS.items():
        caminho = f"{PASTA}\\{gabarito_path}"
        try:
            print(f"\n{gabarito_path}")
            marcar_gabarito(caminho, respostas)
        except Exception as e:
            print(f"  ERRO: {e}")

    print("\n" + "=" * 60)
    print("Concluido!")

if __name__ == "__main__":
    main()
