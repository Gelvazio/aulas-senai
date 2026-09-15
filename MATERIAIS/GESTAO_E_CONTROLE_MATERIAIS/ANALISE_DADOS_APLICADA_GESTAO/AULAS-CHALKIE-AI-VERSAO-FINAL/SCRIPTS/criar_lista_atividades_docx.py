# -*- coding: utf-8 -*-
"""
Criar 1 arquivo DOCX por aula contendo TODAS as atividades daquela aula
Formato: Contexto + Comando para cada atividade
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

PASTA = r"C:\fontes\aulas-senai\MATERIAIS\GESTAO_E_CONTROLE_MATERIAIS\ANALISE_DADOS_APLICADA_GESTAO\AULAS-CHALKIE-AI-VERSAO-FINAL"

ATIVIDADES_PRATICAS = {
    "1-Matematica-Aplicada-a-Gestao": {
        "nome_aula": "Matemática Aplicada à Gestão",
        "atividades": [
            {
                "numero": 1,
                "titulo": "Cálculo de Estoque Final",
                "contexto": "Uma distribuidora de alimentos iniciou a semana com 320 unidades de um produto em estoque. Recebeu 150 unidades novas e vendeu 290 unidades durante a semana.",
                "comando": "Calcule o estoque final da semana utilizando a fórmula: Estoque Final = Estoque Inicial + Compras - Vendas"
            },
            {
                "numero": 2,
                "titulo": "Saldo de Caixa Diário",
                "contexto": "Uma empresa tem saldo inicial de caixa de R$ 5.000. Durante o dia recebeu R$ 3.200 de clientes e pagou R$ 1.800 em despesas operacionais.",
                "comando": "Determine o saldo final do caixa do dia aplicando: Saldo Final = Saldo Inicial + Entradas - Saídas"
            },
            {
                "numero": 3,
                "titulo": "Razão de Produtividade",
                "contexto": "Um departamento de produção tem 5 operários que montam 200 unidades em um turno de 8 horas.",
                "comando": "Calcule a razão de produtividade (unidades por operário) e interprete o resultado"
            },
            {
                "numero": 4,
                "titulo": "Desconto Percentual",
                "contexto": "Um produto custava R$ 200,00 e sofreu um desconto de 15% para uma promoção especial.",
                "comando": "Calcule o preço final do produto após o desconto percentual aplicado"
            },
            {
                "numero": 5,
                "titulo": "Proporção em Receita",
                "contexto": "Uma empresa precisa escalar uma receita de 2 litros de suco para produzir 50 litros, mantendo as mesmas proporções de ingredientes.",
                "comando": "Use proporção (regra de três simples) para calcular quanto de cada ingrediente será necessário"
            },
            {
                "numero": 6,
                "titulo": "Amplitude de Variação",
                "contexto": "Um fornecedor A entrega com variação de 9 a 11 dias e outro fornecedor B entrega com variação de 3 a 19 dias.",
                "comando": "Calcule a amplitude (variação) de cada fornecedor e analise qual é mais confiável para planejamento"
            }
        ]
    },
    "2-Excel-Basico-e-Intermediario": {
        "nome_aula": "Excel Básico e Intermediário para Gestão",
        "atividades": [
            {
                "numero": 1,
                "titulo": "Estrutura da Planilha",
                "contexto": "Você precisa criar uma planilha de vendas mensais com colunas: Produto, Quantidade, Preço Unitário e Total.",
                "comando": "Abra o Excel e estruture a planilha com cabeçalhos formatados em negrito, com bordas e fundo destacado"
            },
            {
                "numero": 2,
                "titulo": "Cálculo de Faturamento",
                "contexto": "Uma loja vendeu: 10 unidades de produto A (R$ 50 cada), 5 unidades de produto B (R$ 80 cada) e 8 unidades de produto C (R$ 30 cada).",
                "comando": "Crie fórmulas para calcular o total de cada produto (Qtd × Preço) e o faturamento total com função SOMA"
            },
            {
                "numero": 3,
                "titulo": "Referência Absoluta",
                "contexto": "Você tem uma lista de preços originais e precisa calcular o preço com desconto de 10% para uma promoção.",
                "comando": "Use referência absoluta ($) para manter fixa a célula de percentual enquanto copia a fórmula para toda coluna"
            },
            {
                "numero": 4,
                "titulo": "Função SE para Bônus",
                "contexto": "Uma empresa quer calcular bônus para vendedores: se vendeu menos de R$ 1.000, não recebe bônus; acima disso, recebe 5%.",
                "comando": "Implemente função SE para avaliar a condição e calcular o bônus de forma automática"
            },
            {
                "numero": 5,
                "titulo": "Contar com CONT.SE",
                "contexto": "Você tem uma lista de 50 clientes com seus status (Ativo, Inativo) e quer contar quantos clientes ativos existem.",
                "comando": "Use função CONT.SE para contar células que atendem ao critério 'Ativo' na coluna de status"
            },
            {
                "numero": 6,
                "titulo": "Soma Condicional SOMASE",
                "contexto": "Você precisa somar as vendas apenas dos produtos da categoria 'Eletrônicos' em uma planilha com muitos produtos.",
                "comando": "Aplique função SOMASE para somar valores onde a coluna categoria seja igual a 'Eletrônicos'"
            }
        ]
    },
    "3-Excel-Avancado-e-Visualizacao": {
        "nome_aula": "Excel Avançado e Visualização de Dados",
        "atividades": [
            {
                "numero": 1,
                "titulo": "PROCV Básico",
                "contexto": "Você tem uma tabela de produtos com código, nome e preço. Precisa buscar o preço de um produto específico pelo código.",
                "comando": "Implemente função PROCV para buscar o preço baseado no código do produto"
            },
            {
                "numero": 2,
                "titulo": "PROCV com Referência Absoluta",
                "contexto": "Você tem duas tabelas: uma com códigos de clientes e nomes, outra com códigos de clientes e valores de compras.",
                "comando": "Use PROCV com referência absoluta para vincular os dados de ambas as tabelas mantendo a estabilidade"
            },
            {
                "numero": 3,
                "titulo": "PROCH para Dados Horizontais",
                "contexto": "Você tem um relatório de vendas com dados em linhas (vendedores como cabeçalhos horizontais) e precisa extrair informações.",
                "comando": "Aplique função PROCH para buscar valores em estrutura horizontal de dados"
            },
            {
                "numero": 4,
                "titulo": "ÍNDICE e CORRESPONDÊNCIA",
                "contexto": "Você precisa buscar um valor que está à ESQUERDA da chave (impossível com PROCV tradicional).",
                "comando": "Combine ÍNDICE e CORRESPONDÊNCIA para fazer buscas flexíveis em qualquer direção"
            },
            {
                "numero": 5,
                "titulo": "Tabela Dinâmica",
                "contexto": "Você tem vendas de múltiplas regiões e precisa resumir por região, período e categoria de produto simultaneamente.",
                "comando": "Crie uma Tabela Dinâmica para sumarizar os dados em múltiplas dimensões e filtrar conforme necessário"
            },
            {
                "numero": 6,
                "titulo": "Segmentadores Interativos",
                "contexto": "Você quer permitir que usuários filtrem um dashboard dinamicamente por período, região e categoria com cliques.",
                "comando": "Adicione Segmentadores (Slicers) conectados às Tabelas Dinâmicas para criar filtros visuais interativos"
            }
        ]
    },
    "4-Dashboards-Executivos": {
        "nome_aula": "Dashboards Executivos e Projeto Final Integrado",
        "atividades": [
            {
                "numero": 1,
                "titulo": "KPIs Críticos para Executivos",
                "contexto": "Uma empresa quer criar um dashboard executivo para a diretoria que resume performance em 10 segundos.",
                "comando": "Defina quais 5 KPIs são críticos para executivos e organize-os no padrão Z (top-left = mais importante)"
            },
            {
                "numero": 2,
                "titulo": "Paleta de Cores Corporativa",
                "contexto": "Você está montando um dashboard com faturamento, margem, quantidade e satisfação de cliente.",
                "comando": "Escolha cores: neutra para fundo, uma cor primária para dados normais, e vermelha APENAS para metas não atingidas"
            },
            {
                "numero": 3,
                "titulo": "Eliminar Poluição Visual",
                "contexto": "Seu dashboard atual tem 15 linhas de grade, sombras 3D, rótulos minúsculos e cores em arco-íris.",
                "comando": "Limpe a poluição visual: remova grades, use gráficos 2D planos, aumente fontes e reduza paleta a 3-4 cores"
            },
            {
                "numero": 4,
                "titulo": "KPI Estruturado",
                "contexto": "Você tem faturamento realizado (R$ 125.500) e meta (R$ 120.000) mas o dashboard apenas mostra o número.",
                "comando": "Estruture o KPI com valor realizado, meta e variação percentual (%), adicionando status visual (✅ ou ⚠️)"
            },
            {
                "numero": 5,
                "titulo": "Variação Percentual",
                "contexto": "Você quer comparar desempenho de vendas com o mesmo período do ano passado.",
                "comando": "Calcule variação percentual: [(Valor Atual - Valor Anterior) / Valor Anterior] × 100"
            },
            {
                "numero": 6,
                "titulo": "Storytelling com Dados",
                "contexto": "Sua equipe tem dificuldade interpretar um dashboard com 8 gráficos diferentes sem contexto de negócio.",
                "comando": "Estruture uma narrativa com dados (storytelling) que explique o problema, a causa e a ação recomendada"
            }
        ]
    }
}

def criar_docx_lista_atividades(nome_arquivo, dados_aula):
    """Criar 1 arquivo DOCX com TODAS as atividades da aula"""

    # Criar documento
    doc = Document()

    # Adicionar título da aula
    titulo_aula = doc.add_paragraph()
    titulo_aula.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo_aula.add_run(f"Lista de Atividades\n{dados_aula['nome_aula']}")
    run.font.size = Pt(16)
    run.font.bold = True

    # Adicionar descrição
    desc = doc.add_paragraph("Formato: Contexto + Comando")
    desc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in desc.runs:
        run.font.italic = True
        run.font.size = Pt(10)

    doc.add_paragraph()  # Espaço

    # Adicionar cada atividade
    for atividade in dados_aula['atividades']:
        # Título da atividade
        titulo_atividade = doc.add_paragraph()
        run = titulo_atividade.add_run(f"Atividade {atividade['numero']}: {atividade['titulo']}")
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

        # Contexto
        contexto_titulo = doc.add_paragraph()
        run = contexto_titulo.add_run("Contexto:")
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 102, 204)  # Azul
        run.font.size = Pt(11)

        contexto_texto = doc.add_paragraph(atividade['contexto'])
        contexto_texto.paragraph_format.left_indent = Inches(0.5)
        contexto_texto.paragraph_format.first_line_indent = Inches(0)

        # Comando
        comando_titulo = doc.add_paragraph()
        run = comando_titulo.add_run("Comando:")
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 102, 204)  # Azul
        run.font.size = Pt(11)

        comando_texto = doc.add_paragraph(atividade['comando'])
        comando_texto.paragraph_format.left_indent = Inches(0.5)
        comando_texto.paragraph_format.first_line_indent = Inches(0)

        # Espaço para resposta
        doc.add_paragraph("_" * 80)
        doc.add_paragraph()  # Espaço em branco

    # Rodapé com instruções
    doc.add_paragraph()
    rodape = doc.add_paragraph("Resolva todas as atividades de forma prática em seu caderno ou planilha. Mostre todas as etapas do seu raciocínio.")
    rodape.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in rodape.runs:
        run.font.size = Pt(9)
        run.font.italic = True

    # Salvar arquivo
    nome_docx = f"ATIVIDADE-{nome_arquivo}.docx"
    caminho = f"{PASTA}\\{nome_docx}"
    doc.save(caminho)

    return nome_docx

def main():
    print("CRIAR ARQUIVO DOCX COM LISTA DE ATIVIDADES POR AULA")
    print("=" * 80)

    for nome_arquivo, dados_aula in ATIVIDADES_PRATICAS.items():
        print(f"\n📚 {dados_aula['nome_aula']}")

        nome_docx = criar_docx_lista_atividades(nome_arquivo, dados_aula)
        num_atividades = len(dados_aula['atividades'])

        print(f"  ✅ {nome_docx}")
        print(f"     {num_atividades} atividades incluídas")

    print("\n" + "=" * 80)
    print("✅ Conclusão!")
    print(f"   Total de arquivos DOCX criados: 4")
    print(f"   Cada arquivo contém uma lista completa de atividades da aula")
    print(f"   Formato: Contexto + Comando em arquivo Word formatado")

if __name__ == "__main__":
    main()
