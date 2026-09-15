# -*- coding: utf-8 -*-
"""
Extrair conteúdo de cada slide PPTX e criar arquivo markdown SEPARADO para cada slide
"""

from pptx import Presentation
import os
from pathlib import Path

PASTA = r"C:\fontes\aulas-senai\MATERIAIS\GESTAO_E_CONTROLE_MATERIAIS\ANALISE_DADOS_APLICADA_GESTAO\AULAS-CHALKIE-AI-VERSAO-FINAL"

# Listar arquivos PPTX
arquivos_pptx = [f for f in os.listdir(PASTA) if f.endswith('.pptx') and f != "APRESENTACAO UNIDADE CURRICULAR.pptx"]
arquivos_pptx.sort()

def extrair_texto_shape(shape):
    """Extrair texto de um shape (caixa de texto, título, etc)"""
    texto = []

    if hasattr(shape, "text_frame"):
        for paragrafo in shape.text_frame.paragraphs:
            if paragrafo.text.strip():
                texto.append(paragrafo.text)

    return "\n".join(texto)

def sanitizar_nome_arquivo(texto):
    """Criar nome de arquivo válido a partir do texto"""
    # Remover caracteres inválidos
    invalidos = r'<>:"/\|?*'
    for char in invalidos:
        texto = texto.replace(char, '')

    # Limitar a 60 caracteres
    texto = texto[:60].strip()

    # Se vazio, retornar genérico
    if not texto:
        return "slide_sem_titulo"

    return texto.lower().replace(' ', '-')

def extrair_slides_pptx(caminho_pptx, prefixo_arquivo):
    """Extrair conteúdo de cada slide e criar arquivo separado"""

    total_criados = 0

    try:
        prs = Presentation(caminho_pptx)

        for slide_idx, slide in enumerate(prs.slides, 1):
            conteudo = []

            # Extrair texto de todas as formas (shapes)
            for shape in slide.shapes:
                texto = extrair_texto_shape(shape)
                if texto.strip():
                    conteudo.append(texto)

            # Preparar conteúdo do markdown
            markdown_content = f"# Slide {slide_idx}\n\n"
            markdown_content += f"**Origem:** {os.path.basename(caminho_pptx)}\n\n"
            markdown_content += "---\n\n"

            if conteudo:
                markdown_content += "\n\n".join(conteudo)
            else:
                markdown_content += "*(Slide sem conteúdo de texto)*"

            # Criar nome do arquivo
            # Usar o primeiro texto como título (primeiros 40 caracteres)
            primeiro_texto = conteudo[0][:40] if conteudo else "slide"
            nome_slide = sanitizar_nome_arquivo(primeiro_texto)
            nome_arquivo = f"{prefixo_arquivo}-{slide_idx:03d}-{nome_slide}.md"

            caminho_saida = os.path.join(PASTA, nome_arquivo)

            # Salvar arquivo
            with open(caminho_saida, 'w', encoding='utf-8') as f:
                f.write(markdown_content)

            total_criados += 1

        return total_criados

    except Exception as e:
        print(f"  ❌ Erro ao ler {caminho_pptx}: {e}")
        return 0

def main():
    print("CRIAR ARQUIVO MARKDOWN SEPARADO PARA CADA SLIDE")
    print("=" * 80)

    total_arquivos = 0

    # Mapear prefixos para cada arquivo
    prefixos = {
        "1-Matemática-Aplicada-à-Gestão.pptx": "01-matematica",
        "2-Excel-Básico-e-Intermediário-para-Gestão.pptx": "02-excel-basico",
        "3-Excel-Avançado-e-Visualização-de-Dados.pptx": "03-excel-avancado",
        "4-Dashboards-Executivos-e-Projeto-Final-Integrado.pptx": "04-dashboards"
    }

    for arquivo in arquivos_pptx:
        caminho = os.path.join(PASTA, arquivo)
        prefixo = prefixos.get(arquivo, "slide")

        print(f"\n📄 {arquivo}")

        total = extrair_slides_pptx(caminho, prefixo)
        total_arquivos += total

        if total > 0:
            print(f"  ✅ {total} arquivos markdown criados")

    print("\n" + "=" * 80)
    print(f"✅ Conclusão!")
    print(f"   Total de arquivos markdown criados: {total_arquivos}")
    print(f"   Localização: {PASTA}")

if __name__ == "__main__":
    main()
