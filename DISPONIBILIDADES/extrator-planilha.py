#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extrator de Resoluções - Cursos Técnicos SENAI
Lê a planilha Excel e extrai dados em JSON estruturado
"""

import openpyxl
import json
from datetime import datetime
import re
from pathlib import Path

# Caminho da planilha
PLANILHA_PATH = r"C:\fontes\aulas-senai\DISPONIBILIDADES\Resoluções_Cursos Técnico.xlsx"
OUTPUT_PATH = r"C:\fontes\aulas-senai\DISPONIBILIDADES\dados-resolucoes.json"

def limpar_valor(valor):
    """Limpa e normaliza valores extraídos"""
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, str):
        return valor.strip() if valor.strip() else None
    return valor

def calcular_status(data_vencimento):
    """Calcula status de validade baseado na data de vencimento"""
    if data_vencimento is None or data_vencimento == "":
        return "desconhecido"

    try:
        if isinstance(data_vencimento, str):
            # Tentar parsear string
            if data_vencimento.startswith("#"):  # Erro Excel
                return "erro"
            data = datetime.strptime(data_vencimento, "%Y-%m-%d").date()
        else:
            data = data_vencimento.date() if hasattr(data_vencimento, 'date') else data_vencimento

        hoje = datetime.now().date()
        dias_faltando = (data - hoje).days

        if dias_faltando < 0:
            return "vencido"
        elif dias_faltando <= 90:
            return "vencendo"
        else:
            return "válido"
    except Exception as e:
        return "erro"

def extrair_dados():
    """Extrai todos os dados da planilha"""
    print("🔄 Iniciando extração de dados...")

    wb = openpyxl.load_workbook(PLANILHA_PATH)
    print(f"✅ Planilha carregada com {len(wb.sheetnames)} abas")

    dados = {
        "meta": {
            "titulo": "Resoluções Cursos Técnicos SENAI",
            "data_extracao": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "total_registros": 0,
            "cursos_unicos": 0
        },
        "cursos": []
    }

    # Abas a pular
    skip_sheets = {"Índice"}

    total_registros = 0
    cursos_processados = set()

    for sheet_name in wb.sheetnames:
        if sheet_name in skip_sheets:
            print(f"⏭️  Pulando aba: {sheet_name}")
            continue

        ws = wb[sheet_name]
        print(f"\n📄 Processando aba: {sheet_name}")

        # Encontrar a linha de cabeçalho (linha 4)
        # Estrutura: linha 1=voltar, linha 2=titulo curso, linha 3=vazio, linha 4=cabeçalho
        header_row = 4
        data_start = 5

        # Extrair nome do curso do título (linha 2)
        titulo_cell = ws[f"A2"].value
        curso_nome = titulo_cell.replace("CT ", "").replace("CT  ", "").strip() if titulo_cell else sheet_name

        # Extrair cabeçalhos
        headers = []
        for col in range(1, 20):
            cell = ws.cell(row=header_row, column=col)
            if cell.value:
                headers.append(cell.value)

        if not headers:
            print(f"  ⚠️  Sem cabeçalhos encontrados, pulando...")
            continue

        # Processar linhas de dados
        resolucoes = []
        for row_num in range(data_start, ws.max_row + 1):
            row_data = []
            tem_dados = False

            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col)
                valor = limpar_valor(cell.value)
                if valor:
                    tem_dados = True
                row_data.append(valor)

            if not tem_dados:
                continue  # Pular linha vazia

            # Montar registro
            registro = {}
            for i, header in enumerate(headers):
                if i < len(row_data):
                    registro[header] = row_data[i]

            # Adicionar campos calculados
            if "Vencimento" in registro:
                vencimento = registro["Vencimento"]
                if isinstance(vencimento, datetime):
                    vencimento_str = vencimento.strftime("%Y-%m-%d")
                else:
                    vencimento_str = str(vencimento) if vencimento else None
                registro["data_vencimento"] = vencimento_str
                registro["status"] = calcular_status(vencimento_str)

            resolucoes.append(registro)
            total_registros += 1

        if resolucoes:
            cursos_processados.add(curso_nome)
            curso_obj = {
                "id": re.sub(r'[^a-z0-9]+', '_', curso_nome.lower()),
                "nome": curso_nome,
                "total_resolucoes": len(resolucoes),
                "resolucoes": resolucoes
            }
            dados["cursos"].append(curso_obj)
            print(f"  ✅ {len(resolucoes)} resoluções extraídas de {curso_nome}")

    # Atualizar meta
    dados["meta"]["total_registros"] = total_registros
    dados["meta"]["cursos_unicos"] = len(cursos_processados)

    print(f"\n{'='*50}")
    print(f"✅ EXTRAÇÃO CONCLUÍDA")
    print(f"   Total de registros: {total_registros}")
    print(f"   Cursos processados: {len(cursos_processados)}")
    print(f"{'='*50}\n")

    wb.close()
    return dados

def salvar_json(dados):
    """Salva dados em JSON"""
    print(f"💾 Salvando JSON em: {OUTPUT_PATH}")

    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)

    print(f"✅ JSON salvo com sucesso!")

    # Mostrar tamanho do arquivo
    tamanho = Path(OUTPUT_PATH).stat().st_size
    print(f"   Tamanho: {tamanho / 1024:.1f} KB")

if __name__ == "__main__":
    try:
        dados = extrair_dados()
        salvar_json(dados)
        print("\n🎉 Tudo pronto! Dashboard pode ser criado agora.")
    except Exception as e:
        print(f"\n❌ ERRO: {e}")
        import traceback
        traceback.print_exc()
